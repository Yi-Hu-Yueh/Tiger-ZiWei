"""DOCX generator, endpoint, and frontend export-state tests."""

from functools import partial
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from fastapi.testclient import TestClient
import pytest

import app.api.report as report_api
from app.birth_input import normalize_birth_input
from app.main import app
from app.models.birth import BirthInput
from app.models.interpretation import InterpretationResult
from app.models.flow_year_interpretation import FlowYearInterpretationResult
from app.models.flow_period_interpretation import (
    FlowDayInterpretationResult,
    FlowMonthInterpretationResult,
)
from app.models.flow_date import FlowDateInput
from app.models.flow_query import FlowQueryInput, LunarFlowQueryInput, SolarFlowQueryInput
from app.models.major_luck_interpretation import MajorLuckInterpretationResult
from app.models.ziwei import PalaceName
from app.report.word_report import (
    DOCX_MEDIA_TYPE,
    DISCLAIMER,
    PALACE_ORDER,
    REPORT_CONTENT_MANIFEST,
    render_word_report,
)
from app.ziwei import calculate_basic_chart
from app.ziwei.flow_year import calculate_flow_year
from app.ziwei.flow_date import calculate_flow_date
from app.ziwei.flow_query import calculate_flow_query


client = TestClient(app)

COMMON = {
    "name": "Case A",
    "gender": "female",
    "birth_hour": 0,
    "birth_minute": 30,
    "birthplace": "Taipei",
}
SOLAR_A = COMMON | {
    "calendar_type": "solar",
    "birth_year": 2025,
    "birth_month": 1,
    "birth_day": 29,
}
SOLAR_L = COMMON | {
    "name": "Case L",
    "calendar_type": "solar",
    "birth_year": 2025,
    "birth_month": 7,
    "birth_day": 25,
    "birth_hour": 12,
    "birth_minute": 0,
}
LUNAR_A = COMMON | {
    "calendar_type": "lunar",
    "lunar_year": 2025,
    "lunar_month": 1,
    "lunar_day": 1,
    "is_leap_month": False,
}


def case_a_interpretation() -> InterpretationResult:
    return InterpretationResult.model_validate(
        {
            "overview": "從傳統紫微斗數的角度，此命盤可作為自我觀察參考。",
            "palace_interpretations": [
                {"palace_name": palace.value, "summary": f"{palace.value}的配置可作為生活反思參考。"}
                for palace in PalaceName
            ],
            "transformation_analysis": "生年四化依命盤既定位置解讀，不改變任何星曜與宮位。",
            "overall": {
                "personality": "性格傾向可從既有配置綜合觀察。",
                "career": "職涯選擇宜結合實際能力與工作環境。",
                "finance": "財務傾向只供傳統命理角度參考。",
                "relationships": "感情互動可以留意溝通與界線。",
                "interpersonal": "人際關係可能重視互信與分工。",
                "family": "家庭議題仍應配合實際生活脈絡理解。",
                "strengths": "優勢在於整合不同資訊並穩定推進。",
                "potential_challenges": "可留意過度承擔與溝通不足。",
            },
        }
    )


def case_a_major_luck_interpretation() -> MajorLuckInterpretationResult:
    return MajorLuckInterpretationResult.model_validate({
        "major_luck_index": 5,
        "start_nominal_age": 45,
        "end_nominal_age": 54,
        "palace_name": "官祿宮",
        "earthly_branch": "午",
        "palace_ganzhi": {"heavenly_stem": "壬", "earthly_branch": "午"},
        "overview": "此大限可留意職涯定位與長期責任。",
        "host_palace_analysis": "大限落本命官祿宮，可觀察工作角色與發展方向。",
        "transformation_analysis": [
            {"transformation_type": "化祿", "star_name": "天梁", "natal_palace_name": "父母宮", "analysis": "天梁化祿可留意經驗與支持資源。"},
            {"transformation_type": "化權", "star_name": "紫微", "natal_palace_name": "官祿宮", "analysis": "紫微化權可能提高承擔與統整要求。"},
            {"transformation_type": "化科", "star_name": "天府", "natal_palace_name": "財帛宮", "analysis": "天府化科適合重視穩健規劃。"},
            {"transformation_type": "化忌", "star_name": "武曲", "natal_palace_name": "命宮", "analysis": "武曲化忌提醒避免過度緊繃。"},
        ],
        "career": "職涯宜建立清楚權責。",
        "finance": "財務宜重視穩健配置。",
        "relationships": "感情可留意工作壓力對溝通的影響。",
        "family_and_interpersonal": "家庭與人際適合重視界線。",
        "strengths": "較能整合資源。",
        "potential_challenges": "可留意責任集中。",
        "practical_focus": "適合建立可持續的工作節奏。",
    })


def case_a_flow_year_interpretation(target_year: int = 2029) -> FlowYearInterpretationResult:
    if target_year == 2026:
        return FlowYearInterpretationResult.model_validate({
            "target_year": 2026,
            "flow_year_ganzhi": {"heavenly_stem": "丙", "earthly_branch": "午"},
            "nominal_age": 2,
            "flow_life_palace_branch": "午",
            "flow_life_palace_natal_host": {"palace_name": "官祿宮", "earthly_branch": "午", "palace_ganzhi": {"heavenly_stem": "壬", "earthly_branch": "午"}},
            "active_major_luck_summary": {"status": "before_first_major_luck", "major_luck_index": None, "start_nominal_age": None, "end_nominal_age": None, "palace_name": None, "earthly_branch": None, "palace_ganzhi": None},
            "overview": "2026 年可由本命與流年層次觀察年度節奏。",
            "flow_life_palace_analysis": "流年命宮落本命官祿宮，可留意日常任務與角色適應。",
            "major_luck_context": "此年尚未進入第一大限，不加入其他運限。",
            "transformation_analysis": [
                {"transformation_type": "化祿", "star_name": "天同", "natal_palace_name": "兄弟宮", "analysis": "天同化祿可觀察互動中的和緩資源。"},
                {"transformation_type": "化權", "star_name": "天機", "natal_palace_name": "田宅宮", "analysis": "天機化權可留意安排與變動。"},
                {"transformation_type": "化科", "star_name": "文昌", "natal_palace_name": "財帛宮", "analysis": "文昌化科可重視清楚表達。"},
                {"transformation_type": "化忌", "star_name": "廉貞", "natal_palace_name": "財帛宮", "analysis": "廉貞化忌提醒維持適當界線。"},
            ],
            "career": "職涯議題以年度觀察為主。", "finance": "財務宜保留彈性。",
            "relationships": "感情互動可重視溝通。", "family_and_interpersonal": "家庭與人際宜維持耐心。",
            "strengths": "能逐步適應環境。", "potential_challenges": "可留意節奏變化。",
            "practical_focus": "以穩定生活節奏為重點。",
        })
    return FlowYearInterpretationResult.model_validate({
        "target_year": 2029,
        "flow_year_ganzhi": {"heavenly_stem": "己", "earthly_branch": "酉"},
        "nominal_age": 5,
        "flow_life_palace_branch": "酉",
        "flow_life_palace_natal_host": {"palace_name": "疾厄宮", "earthly_branch": "酉", "palace_ganzhi": {"heavenly_stem": "乙", "earthly_branch": "酉"}},
        "active_major_luck_summary": {"status": "active", "major_luck_index": 1, "start_nominal_age": 5, "end_nominal_age": 14, "palace_name": "命宮", "earthly_branch": "寅", "palace_ganzhi": {"heavenly_stem": "戊", "earthly_branch": "寅"}},
        "overview": "2029 年可在本命與大限背景下觀察年度重點。",
        "flow_life_palace_analysis": "流年命宮落本命疾厄宮，可留意節奏與負荷。",
        "major_luck_context": "目前為 5 至 14 歲戊寅命宮大限背景。",
        "transformation_analysis": [
            {"transformation_type": "化祿", "star_name": "武曲", "natal_palace_name": "命宮", "analysis": "武曲化祿可留意執行力與資源。"},
            {"transformation_type": "化權", "star_name": "貪狼", "natal_palace_name": "夫妻宮", "analysis": "貪狼化權可留意互動主動性。"},
            {"transformation_type": "化科", "star_name": "天梁", "natal_palace_name": "父母宮", "analysis": "天梁化科可觀察經驗支持。"},
            {"transformation_type": "化忌", "star_name": "文曲", "natal_palace_name": "福德宮", "analysis": "文曲化忌提醒釐清想法。"},
        ],
        "career": "職涯宜按年度節奏推進。", "finance": "財務可重視紀律。",
        "relationships": "感情可留意期待與界線。", "family_and_interpersonal": "家庭與人際宜具體溝通。",
        "strengths": "較能務實整理資源。", "potential_challenges": "可留意壓力集中。",
        "practical_focus": "建立可持續的年度節奏。",
    })


def flow_month_interpretation(result) -> FlowMonthInterpretationResult:
    return FlowMonthInterpretationResult.model_validate({
        "lunar_year": result.lunar_year, "lunar_month": result.lunar_month,
        "is_leap_month": result.is_leap_month,
        "month_ganzhi": result.month_ganzhi.model_dump(mode="json"),
        "flow_month_life_palace_branch": result.flow_month_life_palace_branch,
        "natal_host_palace": result.natal_host_palace_name.value,
        "overview": "流月總覽。", "life_palace_analysis": "流月命宮分析。",
        "major_luck_context": "大限背景。", "flow_year_context": "流年背景。",
        "transformation_analysis": [
            {"transformation_type": item.transformation_type.value, "star_name": item.star_name.value,
             "natal_palace_name": item.natal_palace_name.value, "analysis": "流月四化分析。"}
            for item in result.transformations
        ],
        "career": "流月職涯。", "finance": "流月財務。", "relationships": "流月感情。",
        "family_and_interpersonal": "流月家庭人際。", "strengths": "流月優勢。",
        "potential_challenges": "流月挑戰。", "practical_focus": "流月實際重點。",
    })


def flow_day_interpretation(result) -> FlowDayInterpretationResult:
    return FlowDayInterpretationResult.model_validate({
        "lunar_year": result.target_lunar_year, "lunar_month": result.target_lunar_month,
        "lunar_day": result.target_lunar_day, "is_leap_month": result.target_is_leap_month,
        "day_ganzhi": result.day_ganzhi.model_dump(mode="json"),
        "flow_day_life_palace_branch": result.flow_day_life_palace_branch,
        "natal_host_palace": result.natal_host_palace_name.value,
        "overview": "流日總覽。", "life_palace_analysis": "流日命宮分析。",
        "major_luck_context": "大限背景。", "flow_year_context": "流年背景。",
        "flow_month_context": "流月背景。",
        "transformation_analysis": [
            {"transformation_type": item.transformation_type.value, "star_name": item.star_name.value,
             "natal_palace_name": item.natal_palace_name.value, "analysis": "流日四化分析。"}
            for item in result.transformations
        ],
        "work": "流日工作。", "finance": "流日財務。", "relationships": "流日感情。",
        "family_and_interpersonal": "流日家庭人際。", "strengths": "流日優勢。",
        "potential_challenges": "流日挑戰。", "practical_focus": "流日實際重點。",
    })


def build_report(
    tmp_path: Path,
    payload: dict = SOLAR_A,
    *,
    major_luck_interpretation: MajorLuckInterpretationResult | None = None,
    flow_year: int | None = None,
    flow_date: FlowDateInput | None = None,
    flow_query: FlowQueryInput | None = None,
    flow_year_interpretation: FlowYearInterpretationResult | None = None,
    flow_month_interpretation_result: FlowMonthInterpretationResult | None = None,
    flow_day_interpretation_result: FlowDayInterpretationResult | None = None,
):
    birth_input = BirthInput.model_validate(payload)
    chart = calculate_basic_chart(normalize_birth_input(birth_input))
    flow_date_result = calculate_flow_date(chart, flow_date) if flow_date is not None else None
    flow_query_result = calculate_flow_query(chart, flow_query) if flow_query is not None else None
    return render_word_report(
        chart,
        birth_input,
        case_a_interpretation(),
        flow_date=flow_date_result,
        flow_query=flow_query_result,
        flow_year=(
            flow_query_result.flow_year
            if flow_query_result is not None
            else flow_date_result.flow_year
            if flow_date_result is not None
            else calculate_flow_year(chart, flow_year) if flow_year is not None else None
        ),
        major_luck_interpretation=major_luck_interpretation,
        flow_year_interpretation=flow_year_interpretation,
        flow_month_interpretation=flow_month_interpretation_result,
        flow_day_interpretation=flow_day_interpretation_result,
        output_dir=tmp_path / "reports",
        chart_output_dir=tmp_path / "charts",
    )


def document_text(document: Document) -> str:
    table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*(paragraph.text for paragraph in document.paragraphs), *table_text])


def uncompressed_package_bytes(path: Path) -> bytes:
    with ZipFile(path) as archive:
        return b"\n".join(archive.read(name) for name in archive.namelist())


def report_payload(*, major: bool = False, flow_year: int | None = None) -> dict:
    payload = {
        "birth_input": SOLAR_A,
        "interpretation": case_a_interpretation().model_dump(mode="json"),
    }
    if major:
        payload["major_luck_report"] = {
            "major_luck_index": 5,
            "interpretation": case_a_major_luck_interpretation().model_dump(mode="json"),
        }
    if flow_year is not None:
        payload["target_flow_year"] = flow_year
        payload["flow_year_interpretation"] = case_a_flow_year_interpretation(flow_year).model_dump(mode="json")
    return payload


def flow_date_report_payload(*, leap_target: bool = False, with_interpretation: bool = True) -> dict:
    payload = report_payload(major=True)
    payload["target_datetime"] = (
        {"year": 2025, "month": 7, "day": 25, "hour": 12, "minute": 0}
        if leap_target
        else {"year": 2029, "month": 2, "day": 13, "hour": 12, "minute": 0}
    )
    if with_interpretation and not leap_target:
        payload["flow_year_interpretation"] = case_a_flow_year_interpretation().model_dump(mode="json")
    return payload


def table_with_header(document: Document, header: str):
    return next(table for table in document.tables if table.rows[0].cells[0].text == header)


def test_word_generator_creates_reopenable_structured_docx(tmp_path: Path) -> None:
    rendered = build_report(tmp_path)
    assert rendered.path.is_file()
    assert rendered.path.stat().st_size > 0

    document = Document(rendered.path)
    text = document_text(document)
    assert document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[0].text == "Tiger-ZiWei 紫微斗數完整報告"
    for value in ("姓名", "性別", "輸入方式", "原始出生日期輸入", "換算西元日期", "農曆日期", "出生時間", "出生地"):
        assert value in text
    for heading in (
        "命盤圖片", "本命基本資訊", "本命十二宮表格", "生年四化",
        "本命解讀", "十二宮解讀", "生年四化解讀", "整體分析",
        "大限基本資訊", "十二大限表格", "免責說明",
    ):
        assert heading in text
    assert "AI 模型：GLM-5.3-Flash" in text
    assert "NVIDIA_API_KEY" not in text
    assert "nvapi-" not in text
    assert "reasoning_content" not in text
    assert DISCLAIMER in text


def test_chart_png_is_embedded_before_detailed_interpretation(tmp_path: Path) -> None:
    document = Document(build_report(tmp_path).path)
    image_paragraph_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p.xpath(".//w:drawing")
    )
    detail_index = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "本命基本資訊")
    assert image_paragraph_index < detail_index
    assert len(document.inline_shapes) == 1
    image_relationships = [
        relationship for relationship in document.part.rels.values()
        if relationship.reltype == RELATIONSHIP_TYPE.IMAGE
    ]
    assert len(image_relationships) == 1


def test_report_has_exact_palace_order_facts_and_interpretations(tmp_path: Path) -> None:
    document = Document(build_report(tmp_path).path)
    palace_paragraphs = [
        paragraph.text.split("：", 1)[0]
        for paragraph in document.paragraphs
        if "：" in paragraph.text and paragraph.text.split("：", 1)[0] in {palace.value for palace in PALACE_ORDER}
    ]
    assert palace_paragraphs == [palace.value for palace in PALACE_ORDER]
    text = document_text(document)
    for palace in PALACE_ORDER:
        assert f"{palace.value}的配置可作為生活反思參考。" in text
    for fact in ("武曲", "天相", "陀羅", "天機", "紫微", "太陰", "化祿", "化權", "化科", "化忌"):
        assert fact in text


def test_report_contains_transformation_and_all_overall_fields_only(tmp_path: Path) -> None:
    document = Document(build_report(tmp_path).path)
    text = document_text(document)
    assert case_a_interpretation().transformation_analysis in text
    for value in case_a_interpretation().overall.model_dump().values():
        assert value in text
    assert "大限基本資訊" in text
    assert "十二大限表格" in text
    assert "大限解讀" not in text
    assert "流年基本資訊" not in text
    assert "流年十二宮表格" not in text
    assert "流年解讀" not in text
    assert "reasoning_content" not in text
    assert "NVIDIA_API_KEY" not in text
    package = uncompressed_package_bytes(build_report(tmp_path / "second").path)
    assert b"reasoning_content" not in package
    assert b"NVIDIA_API_KEY" not in package


def test_real_natal_and_deterministic_major_luck_tables_are_complete(tmp_path: Path) -> None:
    document = Document(build_report(tmp_path).path)
    natal = table_with_header(document, "地支")
    assert len(natal.rows) == 13
    chart = calculate_basic_chart(normalize_birth_input(BirthInput.model_validate(SOLAR_A)))
    assert [row.cells[1].text for row in natal.rows[1:]] == [palace.palace_name.value for palace in chart.palaces]

    transformations = table_with_header(document, "四化")
    assert len(transformations.rows) == 5
    assert [row.cells[0].text for row in transformations.rows[1:]] == ["化祿", "化權", "化科", "化忌"]

    major = table_with_header(document, "大限序號")
    assert len(major.rows) == 13
    assert [row.cells[0].text for row in major.rows[1:]] == [str(index) for index in range(1, 13)]
    for row in major.rows[1:]:
        lines = row.cells[5].text.splitlines()
        assert len(lines) == 4
        assert [line.split()[0] for line in lines] == ["祿", "權", "科", "忌"]


def test_dynamic_word_matrix_natal_major_luck_flow_year_and_both(tmp_path: Path) -> None:
    cases = (
        ("natal", None, None, None, False, False, False),
        ("major", case_a_major_luck_interpretation(), None, None, True, False, False),
        ("flow-facts", None, 2029, None, False, True, False),
        ("flow-full", None, 2029, case_a_flow_year_interpretation(), False, True, True),
        ("both", case_a_major_luck_interpretation(), 2029, case_a_flow_year_interpretation(), True, True, True),
    )
    for name, major, flow, flow_interpretation, expect_major, expect_flow, expect_flow_interpretation in cases:
        rendered = build_report(
            tmp_path / name,
            major_luck_interpretation=major,
            flow_year=flow,
            flow_year_interpretation=flow_interpretation,
        )
        document = Document(rendered.path)
        text = document_text(document)
        assert ("大限解讀" in text) is expect_major
        assert ("流年基本資訊" in text) is expect_flow
        assert ("流年解讀" in text) is expect_flow_interpretation
        assert len(document.inline_shapes) == 1
        assert "本命解讀" in text
        assert "十二大限表格" in text
        assert DISCLAIMER in text


def test_dynamic_word_heading_order_and_all_case_a_anchors(tmp_path: Path) -> None:
    rendered = build_report(
        tmp_path,
        major_luck_interpretation=case_a_major_luck_interpretation(),
        flow_year=2029,
        flow_year_interpretation=case_a_flow_year_interpretation(),
    )
    document = Document(rendered.path)
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings == [
        "出生資料", "命盤圖片", "本命基本資訊", "本命十二宮表格",
        "生年四化", "本命解讀", "大限基本資訊", "十二大限表格",
        "大限解讀", "流年基本資訊", "流年十二宮表格", "流年四化",
        "流年解讀", "免責說明",
    ]
    text = document_text(document)
    for expected in (
        "乙巳", "丁丑", "戊戌", "壬子", "命宮", "寅", "土五局",
        "5–14", "戊寅", "15–24", "己卯", "45–54", "官祿宮", "壬午",
        "天梁", "紫微", "天府", "武曲", "2029", "己酉", "疾厄宮｜乙酉",
        "5–14｜戊寅｜命宮", "貪狼", "文曲", "家庭／人際", "實際重點",
    ):
        assert expected in text
    assert len(table_with_header(document, "地支").rows) == 13
    assert len(table_with_header(document, "大限序號").rows) == 13
    assert len(table_with_header(document, "流年宮位").rows) == 13
    assert headings.count("本命十二宮表格") == 1
    assert headings.count("十二大限表格") == 1
    assert headings.count("流年十二宮表格") == 1
    assert headings.count("大限解讀") == 1
    assert headings.count("流年解讀") == 1
    assert headings.count("免責說明") == 1
    assert "NVIDIA_API_KEY" not in text
    assert "reasoning_content" not in text
    package = uncompressed_package_bytes(rendered.path)
    assert b"NVIDIA_API_KEY" not in package
    assert b"reasoning_content" not in package


def test_before_first_major_luck_flow_year_word_is_honest(tmp_path: Path) -> None:
    document = Document(build_report(
        tmp_path,
        flow_year=2026,
        flow_year_interpretation=case_a_flow_year_interpretation(2026),
    ).path)
    text = document_text(document)
    assert "尚未進入第一大限" in text
    assert "2026" in text
    assert "丙午" in text
    assert "小限" not in text
    assert "童限" not in text


def test_filename_sanitizes_windows_characters_and_supports_unnamed_input(tmp_path: Path) -> None:
    invalid_name = SOLAR_A | {"name": 'A<B>C:D"E/F\\G|H?I*J'}
    named = build_report(tmp_path / "named", invalid_name)
    assert not any(character in named.filename for character in '<>:"/\\|?*')
    assert named.filename.endswith("_紫微斗數命盤.docx")

    unnamed = build_report(tmp_path / "unnamed", SOLAR_A | {"name": None})
    assert unnamed.filename == "2025-01-29_0030_紫微斗數命盤.docx"


def test_gregorian_and_lunar_reports_share_identical_chart_and_report_body(tmp_path: Path) -> None:
    solar = build_report(tmp_path / "solar", SOLAR_A)
    lunar = build_report(tmp_path / "lunar", LUNAR_A)
    solar_doc = Document(solar.path)
    lunar_doc = Document(lunar.path)

    def body_paragraphs(document: Document) -> list[str]:
        start = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "本命基本資訊")
        return [paragraph.text for paragraph in document.paragraphs[start:]]

    assert body_paragraphs(solar_doc) == body_paragraphs(lunar_doc)
    assert [[cell.text for row in table.rows for cell in row.cells] for table in solar_doc.tables[1:]] == [
        [cell.text for row in table.rows for cell in row.cells] for table in lunar_doc.tables[1:]
    ]
    with ZipFile(solar.path) as solar_zip, ZipFile(lunar.path) as lunar_zip:
        solar_image = solar_zip.read(next(name for name in solar_zip.namelist() if name.startswith("word/media/")))
        lunar_image = lunar_zip.read(next(name for name in lunar_zip.namelist() if name.startswith("word/media/")))
    assert solar_image == lunar_image
    assert "國曆" in document_text(solar_doc)
    assert "農曆" in document_text(lunar_doc)


def test_report_endpoint_returns_real_docx_without_nvidia_call(tmp_path: Path, monkeypatch) -> None:
    original = report_api.render_word_report
    monkeypatch.setattr(
        report_api,
        "render_word_report",
        partial(original, output_dir=tmp_path / "reports", chart_output_dir=tmp_path / "charts"),
    )

    async def forbidden(*args, **kwargs):
        raise AssertionError("NVIDIA/interpreter must not run during DOCX generation")

    monkeypatch.setattr("app.llm.interpreter.ZiweiInterpreter.interpret", forbidden)
    monkeypatch.setattr("app.llm.major_luck_interpreter.MajorLuckInterpreter.interpret", forbidden)
    monkeypatch.setattr("app.llm.flow_year_interpreter.FlowYearInterpreter.interpret", forbidden)
    monkeypatch.setattr("app.llm.nvidia_client.NvidiaClient.chat", forbidden)
    response = client.post(
        "/api/report/docx",
        json={"birth_input": SOLAR_A, "interpretation": case_a_interpretation().model_dump(mode="json")},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert "content-disposition" in response.headers
    assert ".docx" in response.headers["content-disposition"]
    downloaded = tmp_path / "downloaded.docx"
    downloaded.write_bytes(response.content)
    assert Document(downloaded).paragraphs[0].text == "Tiger-ZiWei 紫微斗數完整報告"


def test_report_endpoint_includes_deterministic_flow_year_without_ai_layer(tmp_path: Path, monkeypatch) -> None:
    original = report_api.render_word_report
    monkeypatch.setattr(
        report_api,
        "render_word_report",
        partial(original, output_dir=tmp_path / "reports", chart_output_dir=tmp_path / "charts"),
    )
    response = client.post(
        "/api/report/docx",
        json=report_payload() | {"target_flow_year": 2029},
    )
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    text = document_text(document)
    assert "流年基本資訊" in text
    assert "流年十二宮表格" in text
    assert "流年四化" in text
    assert "流年解讀" not in text
    assert len(table_with_header(document, "流年宮位").rows) == 13


def test_report_endpoint_authoritatively_rebuilds_flow_year(tmp_path: Path, monkeypatch) -> None:
    original_render = report_api.render_word_report
    original_calculate = report_api.calculate_flow_year
    requested_years: list[int] = []

    def tracked_calculate(chart, target_year):
        requested_years.append(target_year)
        return original_calculate(chart, target_year)

    monkeypatch.setattr(report_api, "calculate_flow_year", tracked_calculate)
    monkeypatch.setattr(
        report_api,
        "render_word_report",
        partial(original_render, output_dir=tmp_path / "reports", chart_output_dir=tmp_path / "charts"),
    )
    response = client.post(
        "/api/report/docx",
        json=report_payload() | {"target_flow_year": 2029},
    )
    assert response.status_code == 200
    assert requested_years == [2029]


def test_full_flow_date_word_contains_all_new_deterministic_content(tmp_path: Path) -> None:
    target = FlowDateInput(year=2029, month=2, day=13, hour=12, minute=0)
    rendered = build_report(
        tmp_path,
        major_luck_interpretation=case_a_major_luck_interpretation(),
        flow_date=target,
        flow_year_interpretation=case_a_flow_year_interpretation(),
    )
    assert rendered.filename == "Case A_紫微斗數命盤_流年2029-流月1-流日1.docx"
    document = Document(rendered.path)
    text = document_text(document)
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings[-12:] == [
        "運限查詢時間", "流年基本資訊", "流年十二宮表格", "流年四化",
        "流年解讀", "流月基本資訊", "流月十二宮表格", "流月四化",
        "流日基本資訊", "流日十二宮表格", "流日四化", "免責說明",
    ]
    for expected in (
        "2029-02-13 12:00", "2029 年正月初一", "流月基本資訊", "流月命宮",
        "流日基本資訊", "流日命宮", "流年基本資訊", "十二大限表格", "本命解讀",
    ):
        assert expected in text
    assert len(table_with_header(document, "流月宮位").rows) == 13
    assert len(table_with_header(document, "流日宮位").rows) == 13
    assert len(document.inline_shapes) == 1
    assert headings.count("流月十二宮表格") == 1
    assert headings.count("流日十二宮表格") == 1


def test_flow_date_word_filename_uses_leap_marker_and_no_target_fallback_is_unchanged(tmp_path: Path) -> None:
    leap = build_report(
        tmp_path / "leap",
        flow_date=FlowDateInput(year=2025, month=7, day=25, hour=12, minute=0),
    )
    assert leap.filename == "Case A_紫微斗數命盤_流年2025-流月閏6-流日1.docx"
    assert build_report(tmp_path / "fallback").filename == "Case A_紫微斗數命盤.docx"


@pytest.mark.parametrize(
    ("query", "filename", "included", "excluded"),
    (
        (
            SolarFlowQueryInput(year=2029, month=2, day=13),
            "Case A_紫微斗數命盤_流年2029-流月1-流日1.docx",
            ("查詢類型", "國曆", "輸入日期", "2029-02-13", "轉換農曆", "2029 年正月初一", "流月基本資訊", "流日基本資訊"),
            (),
        ),
        (
            LunarFlowQueryInput(year=2026),
            "Case A_紫微斗數命盤_流年2026.docx",
            ("查詢類型", "農曆", "輸入", "2026", "流年基本資訊"),
            ("流月基本資訊", "流日基本資訊"),
        ),
        (
            LunarFlowQueryInput(year=2026, month=10),
            "Case A_紫微斗數命盤_流年2026-流月10.docx",
            ("2026 十月", "流月基本資訊"),
            ("流日基本資訊",),
        ),
        (
            LunarFlowQueryInput(year=2026, month=10, day=7),
            "Case A_紫微斗數命盤_流年2026-流月10-流日7.docx",
            ("2026 十月初七", "流月基本資訊", "流日基本資訊"),
            (),
        ),
        (
            LunarFlowQueryInput(year=2025, month=6, is_leap_month=True),
            "Case A_紫微斗數命盤_流年2025-流月閏6.docx",
            ("2025 閏六月", "有效月份", "7"),
            ("流日基本資訊",),
        ),
        (
            LunarFlowQueryInput(year=2025, month=6, day=1, is_leap_month=True),
            "Case A_紫微斗數命盤_流年2025-流月閏6-流日1.docx",
            ("2025 閏六月初一", "流月基本資訊", "流日基本資訊"),
            (),
        ),
    ),
)
def test_flow_query_word_filename_metadata_and_granularity(
    tmp_path: Path,
    query: FlowQueryInput,
    filename: str,
    included: tuple[str, ...],
    excluded: tuple[str, ...],
) -> None:
    rendered = build_report(tmp_path / filename, flow_query=query)
    assert rendered.filename == filename
    text = document_text(Document(rendered.path))
    assert "運限查詢" in text
    for expected in included:
        assert expected in text
    for forbidden in excluded:
        assert forbidden not in text


def test_gregorian_leap_query_word_uses_normalized_lunar_filename(tmp_path: Path) -> None:
    rendered = build_report(
        tmp_path,
        SOLAR_L,
        flow_query=SolarFlowQueryInput(year=2025, month=7, day=25),
    )
    assert rendered.filename == "Case L_紫微斗數命盤_流年2025-流月閏6-流日1.docx"
    text = document_text(Document(rendered.path))
    assert "2025-07-25" in text
    assert "2025 年閏六月初一" in text


def test_word_includes_month_day_transformations_and_existing_ai_results_in_layer_order(
    tmp_path: Path,
) -> None:
    birth_input = BirthInput.model_validate(SOLAR_A)
    chart = calculate_basic_chart(normalize_birth_input(birth_input))
    query = SolarFlowQueryInput(year=2029, month=2, day=13)
    result = calculate_flow_query(chart, query)
    rendered = build_report(
        tmp_path,
        flow_query=query,
        flow_year_interpretation=case_a_flow_year_interpretation(),
        flow_month_interpretation_result=flow_month_interpretation(result.flow_month),
        flow_day_interpretation_result=flow_day_interpretation(result.flow_day),
    )
    document = Document(rendered.path)
    text = document_text(document)
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    ordered = [
        "流年基本資訊", "流年十二宮表格", "流年四化", "流年解讀",
        "流月基本資訊", "流月十二宮表格", "流月四化", "流月解讀",
        "流日基本資訊", "流日十二宮表格", "流日四化", "流日解讀",
        "免責說明",
    ]
    assert [heading for heading in headings if heading in ordered] == ordered
    for expected in (
        "天同", "天機", "文昌", "廉貞", "破軍", "武曲", "太陽",
        "流月總覽。", "流月四化分析。", "流日總覽。", "流日四化分析。",
    ):
        assert expected in text
    assert len(table_with_header(document, "四化").rows) == 5


def test_report_endpoint_fact_locks_month_day_interpretations_without_nvidia(
    tmp_path: Path, monkeypatch
) -> None:
    original = report_api.render_word_report
    monkeypatch.setattr(
        report_api,
        "render_word_report",
        partial(original, output_dir=tmp_path / "reports", chart_output_dir=tmp_path / "charts"),
    )

    async def forbidden(*args, **kwargs):
        raise AssertionError("Word generation must not call NVIDIA")

    monkeypatch.setattr("app.llm.nvidia_client.NvidiaClient.chat", forbidden)
    chart = calculate_basic_chart(normalize_birth_input(BirthInput.model_validate(SOLAR_A)))
    query = SolarFlowQueryInput(year=2029, month=2, day=13)
    result = calculate_flow_query(chart, query)
    payload = report_payload()
    payload["flow_query"] = query.model_dump(mode="json")
    payload["flow_month_interpretation"] = flow_month_interpretation(result.flow_month).model_dump(mode="json")
    payload["flow_day_interpretation"] = flow_day_interpretation(result.flow_day).model_dump(mode="json")
    response = client.post("/api/report/docx", json=payload)
    assert response.status_code == 200
    text = document_text(Document(BytesIO(response.content)))
    assert "流月解讀" in text and "流日解讀" in text
    assert b"NVIDIA_API_KEY" not in response.content
    assert b"reasoning_content" not in response.content

    invalid = dict(payload)
    invalid["flow_day_interpretation"] = deepcopy(payload["flow_day_interpretation"])
    invalid["flow_day_interpretation"]["lunar_day"] = 2
    rejected = client.post("/api/report/docx", json=invalid)
    assert rejected.status_code == 422
    assert rejected.json()["detail"]["code"] == "INVALID_FLOW_DAY_REPORT"


def test_report_endpoint_rebuilds_flow_query_without_nvidia(tmp_path: Path, monkeypatch) -> None:
    original = report_api.render_word_report
    monkeypatch.setattr(
        report_api,
        "render_word_report",
        partial(original, output_dir=tmp_path / "reports", chart_output_dir=tmp_path / "charts"),
    )

    async def forbidden(*args, **kwargs):
        raise AssertionError("Word generation must not call NVIDIA")

    monkeypatch.setattr("app.llm.nvidia_client.NvidiaClient.chat", forbidden)
    payload = report_payload()
    payload["flow_query"] = {"mode": "solar", "year": 2029, "month": 2, "day": 13}
    response = client.post("/api/report/docx", json=payload)
    assert response.status_code == 200
    assert "流年2029-流月1-流日1.docx" in unquote(response.headers["content-disposition"])
    text = document_text(Document(BytesIO(response.content)))
    assert "查詢類型" in text and "國曆" in text
    assert "輸入日期" in text and "2029-02-13" in text
    assert "轉換農曆" in text and "2029 年正月初一" in text
    assert b"NVIDIA_API_KEY" not in response.content


def test_report_endpoint_rebuilds_flow_date_without_nvidia(tmp_path: Path, monkeypatch) -> None:
    original = report_api.render_word_report
    monkeypatch.setattr(
        report_api,
        "render_word_report",
        partial(original, output_dir=tmp_path / "reports", chart_output_dir=tmp_path / "charts"),
    )

    async def forbidden(*args, **kwargs):
        raise AssertionError("Word generation must not call NVIDIA")

    monkeypatch.setattr("app.llm.nvidia_client.NvidiaClient.chat", forbidden)
    response = client.post("/api/report/docx", json=flow_date_report_payload())
    assert response.status_code == 200
    assert "流年2029-流月1-流日1.docx" in unquote(response.headers["content-disposition"])
    document = Document(BytesIO(response.content))
    assert len(table_with_header(document, "流月宮位").rows) == 13
    assert len(table_with_header(document, "流日宮位").rows) == 13
    package = response.content
    assert b"NVIDIA_API_KEY" not in package
    assert b"reasoning_content" not in package


def test_report_endpoint_revalidates_and_includes_both_optional_layers_without_nvidia(
    tmp_path: Path,
    monkeypatch,
) -> None:
    original = report_api.render_word_report
    monkeypatch.setattr(
        report_api,
        "render_word_report",
        partial(original, output_dir=tmp_path / "reports", chart_output_dir=tmp_path / "charts"),
    )

    async def forbidden(*args, **kwargs):
        raise AssertionError("no interpreter or NVIDIA call is allowed during DOCX generation")

    monkeypatch.setattr("app.llm.interpreter.ZiweiInterpreter.interpret", forbidden)
    monkeypatch.setattr("app.llm.major_luck_interpreter.MajorLuckInterpreter.interpret", forbidden)
    monkeypatch.setattr("app.llm.flow_year_interpreter.FlowYearInterpreter.interpret", forbidden)
    monkeypatch.setattr("app.llm.nvidia_client.NvidiaClient.chat", forbidden)
    response = client.post("/api/report/docx", json=report_payload(major=True, flow_year=2029))
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    text = document_text(document)
    assert "大限解讀" in text
    assert "流年解讀" in text
    assert "45–54" in text and "5–14｜戊寅｜命宮" in text
    assert len(document.inline_shapes) == 1


def test_report_endpoint_rejects_corrupt_or_stale_optional_anchors() -> None:
    corrupt_major = report_payload(major=True)
    corrupt_major["major_luck_report"]["interpretation"]["transformation_analysis"][2]["star_name"] = "左輔"
    response = client.post("/api/report/docx", json=corrupt_major)
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "INVALID_MAJOR_LUCK_REPORT"

    stale_major = report_payload(major=True)
    stale_major["major_luck_report"]["major_luck_index"] = 4
    response = client.post("/api/report/docx", json=stale_major)
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "INVALID_MAJOR_LUCK_REPORT"

    corrupt_flow = report_payload(flow_year=2029)
    corrupt_flow["flow_year_interpretation"]["transformation_analysis"][0]["star_name"] = "太陽"
    response = client.post("/api/report/docx", json=corrupt_flow)
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "INVALID_FLOW_YEAR_REPORT"

    stale_flow = report_payload(flow_year=2029)
    stale_flow["target_flow_year"] = 2030
    response = client.post("/api/report/docx", json=stale_flow)
    assert response.status_code == 422

    missing_flow = report_payload(flow_year=2029)
    del missing_flow["target_flow_year"]
    response = client.post("/api/report/docx", json=missing_flow)
    assert response.status_code == 422


def test_report_endpoint_rejects_corrupt_natal_interpretation() -> None:
    payload = report_payload()
    payload["interpretation"]["palace_interpretations"][0]["palace_name"] = "兄弟宮"
    response = client.post("/api/report/docx", json=payload)
    assert response.status_code == 422


def test_report_endpoint_accepts_lunar_input(tmp_path: Path, monkeypatch) -> None:
    original = report_api.render_word_report
    monkeypatch.setattr(
        report_api,
        "render_word_report",
        partial(original, output_dir=tmp_path / "reports", chart_output_dir=tmp_path / "charts"),
    )
    response = client.post(
        "/api/report/docx",
        json={"birth_input": LUNAR_A, "interpretation": case_a_interpretation().model_dump(mode="json")},
    )
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    text = document_text(document)
    assert "農曆" in text
    assert "2025 年正月初一" in text
    assert "2025-01-29" in text


def test_ui_word_button_state_download_and_stale_safety() -> None:
    html = client.get("/ziwei").text
    javascript = client.get("/static/ziwei.js").text
    assert 'id="report-button" type="button" class="secondary" disabled' in html
    assert "下載 Word 報告" in html
    assert 'id="interpretation-heading">本命解讀<' in html
    assert 'id="major-luck-interpretation-heading">大限解讀<' in html
    assert 'id="flow-year-interpretation-heading">流年解讀<' in html
    assert "正在產生 Word 報告……" in javascript
    assert javascript.count('fetch("/api/report/docx"') == 1
    assert "currentInterpretation = payload" in javascript
    assert "reportButton.disabled = false" in javascript
    assert "currentInterpretation = null" in javascript
    assert "currentChart = null" in javascript
    assert "currentMajorLuckInterpretation = null" in javascript
    assert "currentFlowYearResult = null" in javascript
    assert "currentFlowYearInterpretation = null" in javascript
    assert "request.major_luck_report" in javascript
    assert "request.flow_query" in javascript
    assert "request.flow_year_interpretation" in javascript
    assert "request.flow_year_report" not in javascript
    assert "currentMajorLuckInterpretationIndex === selectedMajorLuckIndex" in javascript
    assert "currentFlowYearTarget === currentFlowYearResult.target_lunar_year" in javascript
    assert "currentFlowYearTarget === currentFlowYearInterpretation.target_year" in javascript
    assert "reportButton.disabled = true" in javascript
    assert 'form.addEventListener("input"' in javascript
    assert "invalidateInterpretation();" in javascript
    assert 'fetch("/api/interpret"' not in javascript[javascript.index('reportButton.addEventListener("click"'):]
    report_handler = javascript[javascript.index('reportButton.addEventListener("click"'):]
    assert 'fetch("/api/major-luck/interpret"' not in report_handler
    assert 'fetch("/api/flow-year/interpret"' not in report_handler
    assert javascript.count('fetch("/api/report/docx"') == 1


def test_web_to_word_coverage_manifest_is_explicit_and_complete() -> None:
    assert REPORT_CONTENT_MANIFEST == (
        ("出生資料", "出生資料"),
        ("本命命盤圖片", "命盤圖片"),
        ("基本命盤資訊", "本命基本資訊"),
        ("十二宮表格", "本命十二宮表格"),
        ("生年四化", "生年四化"),
        ("本命解讀", "本命解讀"),
        ("大限摘要", "大限基本資訊"),
        ("12-row 大限 table", "十二大限表格"),
        ("大限四化 column/content", "十二大限四化 included in the table"),
        ("selected 大限解讀", "大限解讀"),
        ("流年摘要", "流年基本資訊"),
        ("12-row 流年 table", "流年十二宮表格"),
        ("流年四化", "流年四化"),
        ("流年解讀", "流年解讀"),
        ("運限查詢", "運限查詢"),
        ("流月摘要", "流月基本資訊"),
        ("流月十二宮", "流月十二宮表格"),
        ("流月四化", "流月四化"),
        ("流月解讀", "流月解讀"),
        ("流日摘要", "流日基本資訊"),
        ("流日十二宮", "流日十二宮表格"),
        ("流日四化", "流日四化"),
        ("流日解讀", "流日解讀"),
    )


def test_ui_report_payload_excludes_stale_dynamic_state() -> None:
    javascript = client.get("/static/ziwei.js").text
    major_change = javascript[javascript.index('majorLuckSelect.addEventListener("change"'):]
    major_change = major_change[:major_change.index("});") + 3]
    assert "selectedMajorLuckIndex" in major_change
    assert "invalidateMajorLuckInterpretation();" in major_change

    flow_change = javascript[javascript.index("flowQueryInputs.forEach"):]
    flow_change = flow_change[:flow_change.index("});") + 3]
    assert "invalidateFlowYear();" in flow_change

    birth_change = javascript[javascript.index('form.addEventListener("input"'):]
    birth_change = birth_change[:birth_change.index("});") + 3]
    assert "invalidateInterpretation();" in birth_change
    assert "reportButton.disabled = !isStandaloneAndroid || !currentBirthInput" in javascript[javascript.index("function clearInterpretation()"):javascript.index("function invalidateInterpretation()")]

    report_handler = javascript[javascript.index('reportButton.addEventListener("click"'):]
    assert "currentMajorLuckInterpretationIndex === selectedMajorLuckIndex" in report_handler
    assert "currentFlowYearTarget === currentFlowYearResult.target_lunar_year" in report_handler
    assert "request.flow_query = { ...currentFlowQueryInput }" in report_handler
    assert "currentFlowYearTarget === currentFlowYearInterpretation.target_year" in report_handler
