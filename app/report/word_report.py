"""Generate a readable DOCX from one authoritative chart and interpretation."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from app.chart import render_chart_png
from app.llm.model_registry import default_model, get_model_config
from app.models.basic_chart import BasicChartResult
from app.models.birth import BirthInput, CalendarType
from app.models.flow_year import FlowYearResult
from app.models.flow_date import FlowDateResult, FlowDayResult, FlowMonthResult
from app.models.flow_query import FlowQueryResult, LunarFlowQueryInput, SolarFlowQueryInput
from app.models.flow_year_interpretation import ActiveMajorLuckStatus, FlowYearInterpretationResult
from app.models.flow_period_interpretation import (
    FlowDayInterpretationResult,
    FlowMonthInterpretationResult,
)
from app.models.interpretation import InterpretationResult
from app.models.major_luck_interpretation import MajorLuckInterpretationResult
from app.models.ziwei import PalaceName


DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DEFAULT_REPORT_DIR = Path(__file__).resolve().parents[2] / "output" / "reports"
PALACE_ORDER = tuple(PalaceName)
DISCLAIMER = (
    "本報告內容屬傳統紫微斗數觀點，供文化研究與自我反思參考，"
    "不代表科學結論，也不應取代醫療、法律或財務等專業建議。"
)
REPORT_CONTENT_MANIFEST: tuple[tuple[str, str], ...] = (
    ("出生資料", "出生資料"),
    ("本命命盤圖片", "命盤圖片"),
    ("基本命盤資訊", "本命基本資訊"),
    ("十二宮表格", "本命十二宮表格"),
    ("生年四化", "生年四化"),
    ("本命解讀", "本命解讀"),
    ("大限摘要", "大限基本資訊"),
    ("12-row 大限 table", "十二大限表格"),
    ("大限四化 column/content", "十二大限四化 included in the table"),
    ("selected 大限解讀", "大限解讀"),
    ("流年摘要", "流年基本資訊"),
    ("12-row 流年 table", "流年十二宮表格"),
    ("流年四化", "流年四化"),
    ("流年解讀", "流年解讀"),
    ("運限查詢", "運限查詢"),
    ("流月摘要", "流月基本資訊"),
    ("流月十二宮", "流月十二宮表格"),
    ("流月四化", "流月四化"),
    ("流月解讀", "流月解讀"),
    ("流日摘要", "流日基本資訊"),
    ("流日十二宮", "流日十二宮表格"),
    ("流日四化", "流日四化"),
    ("流日解讀", "流日解讀"),
)
ADOPTED_RULES: tuple[str, ...] = (
    "年界：農曆正月初一，不以立春為界。",
    "時間：民用／標準時間；真太陽時尚未實作。",
    "晚子時：23:xx 不換日。",
    "閏月：整個閏月按次月處理本命月規則。",
    "壬年四化：天梁化祿、紫微化權、天府化科、武曲化忌。",
    "大限：命宮為第一限，五行局數為起限虛歲；陽男陰女順行、陰男陽女逆行，每限十個虛歲。",
    "流年：使用明確輸入的農曆年，以正月初一為年界；流年命宮落在該年地支。",
)


@dataclass(frozen=True, slots=True)
class RenderedWordReport:
    path: Path
    filename: str


def _set_run_font(run, size: float | None = None, *, bold: bool | None = None) -> None:
    run.font.name = "Microsoft JhengHei"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft JhengHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _set_style_font(style, size: float, *, bold: bool = False) -> None:
    style.font.name = "Microsoft JhengHei"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft JhengHei")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), "D9D9D9")


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, value: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _format_table(
    table,
    *,
    label_column: bool = False,
    header: bool = False,
    font_size: float = 10.5,
) -> None:
    table.autofit = True
    _set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        properties.append(cant_split)
        if header and row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            properties.append(repeat)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if header and row_index == 0:
                _shade_cell(cell, "4F5F73")
            elif label_column and column_index == 0:
                _shade_cell(cell, "EAF0F6")
            elif header and row_index % 2 == 0:
                _shade_cell(cell, "F4F7FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    _set_run_font(run, font_size, bold=(header and row_index == 0) or (label_column and column_index == 0))
                    if header and row_index == 0:
                        run.font.color.rgb = RGBColor(255, 255, 255)


def _set_table_widths(table, widths: tuple[float, ...]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Cm(width)


def _month_text(month: int) -> str:
    return ("", "正", "二", "三", "四", "五", "六", "七", "八", "九", "十", "冬", "臘")[month]


def _day_text(day: int) -> str:
    return (
        "", "初一", "初二", "初三", "初四", "初五", "初六", "初七", "初八", "初九", "初十",
        "十一", "十二", "十三", "十四", "十五", "十六", "十七", "十八", "十九", "二十",
        "廿一", "廿二", "廿三", "廿四", "廿五", "廿六", "廿七", "廿八", "廿九", "三十",
    )[day]


def _lunar_text(year: int, month: int, day: int, leap: bool) -> str:
    return f"{year} 年{'閏' if leap else ''}{_month_text(month)}月{_day_text(day)}"


def _original_input_text(value: BirthInput) -> str:
    if value.calendar_type is CalendarType.LUNAR:
        return _lunar_text(value.lunar_year, value.lunar_month, value.lunar_day, value.is_leap_month)
    return f"{value.birth_year:04d}-{value.birth_month:02d}-{value.birth_day:02d}"


def _safe_filename(
    chart: BasicChartResult,
    flow_date: FlowDateResult | None = None,
    flow_query: FlowQueryResult | None = None,
) -> str:
    birth = chart.birth_data
    if birth.name:
        base = birth.name.strip()
    else:
        base = (
            f"{birth.birth_year:04d}-{birth.birth_month:02d}-{birth.birth_day:02d}_"
            f"{birth.birth_hour:02d}{birth.birth_minute:02d}"
        )
    base = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", base).strip(" .")
    suffix = ""
    if flow_query is not None:
        target = flow_query.normalized_target
        suffix = f"_流年{target.lunar_year}"
        if target.lunar_month is not None:
            month = f"{'閏' if target.is_leap_month else ''}{target.lunar_month}"
            suffix += f"-流月{month}"
        if target.lunar_day is not None:
            suffix += f"-流日{target.lunar_day}"
    elif flow_date is not None:
        month = f"{'閏' if flow_date.target_is_leap_month else ''}{flow_date.target_lunar_month}"
        suffix = f"_流年{flow_date.target_lunar_year}-流月{month}-流日{flow_date.target_lunar_day}"
    return f"{(base or 'Tiger-ZiWei')[:120]}_紫微斗數命盤{suffix}.docx"


def _add_key_value_table(document: Document, rows: tuple[tuple[str, str], ...]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    _format_table(table, label_column=True)


def _add_fact_paragraph(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    label_run = paragraph.add_run(f"{label}：")
    _set_run_font(label_run, 10.5, bold=True)
    _set_run_font(paragraph.add_run(value), 10.5)


def _add_ai_model_metadata(document: Document, interpretation) -> None:
    """Identify the trusted model without ever including credentials."""

    display_name = interpretation.model_display_name
    if display_name is None and interpretation.model is not None:
        try:
            display_name = get_model_config(interpretation.model).display_name
        except ValueError:
            display_name = interpretation.model
    document.add_paragraph(f"AI 模型：{display_name or default_model().display_name}")


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    _set_style_font(document.styles["Normal"], 11)
    _set_style_font(document.styles["Title"], 22, bold=True)
    _set_style_font(document.styles["Heading 1"], 16, bold=True)
    _set_style_font(document.styles["Heading 2"], 13, bold=True)
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    document.styles["Normal"].paragraph_format.line_spacing = 1.2
    document.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    document.styles["Heading 1"].paragraph_format.keep_with_next = True
    document.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(5)
    document.styles["Heading 2"].paragraph_format.keep_with_next = True


def _add_analysis_sections(document: Document, values: tuple[tuple[str, str], ...]) -> None:
    for label, text in values:
        document.add_paragraph(label, style="Heading 2")
        document.add_paragraph(text)


def _add_natal_palace_table(document: Document, chart: BasicChartResult) -> None:
    document.add_paragraph("本命十二宮表格", style="Heading 1")
    transformations_by_branch: dict[str, list[str]] = {}
    for item in chart.birth_year_transformations.transformations:
        transformations_by_branch.setdefault(item.earthly_branch, []).append(
            f"{item.transformation.value.replace('化', '')} {item.star_name.value}"
        )
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ("地支", "宮名", "宮干支", "命宮", "身宮", "主星", "輔星", "生年四化")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
    for palace in chart.palaces:
        cells = table.add_row().cells
        values = (
            palace.earthly_branch,
            palace.palace_name.value,
            palace.palace_ganzhi.display,
            "命" if palace.is_life_palace else "",
            "身" if palace.has_body_palace else "",
            "、".join(star.name.value for star in palace.major_stars),
            "、".join(star.name.value for star in palace.auxiliary_stars),
            "\n".join(transformations_by_branch.get(palace.earthly_branch, ())),
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(table, (1.0, 1.6, 1.4, 0.8, 0.8, 2.5, 3.8, 3.5))
    _format_table(table, header=True, font_size=9)


def _add_birth_year_transformations(document: Document, chart: BasicChartResult) -> None:
    document.add_paragraph("生年四化", style="Heading 1")
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("四化", "星曜", "類型", "地支", "宮位"), strict=True):
        cell.text = text
    for item in chart.birth_year_transformations.transformations:
        cells = table.add_row().cells
        values = (
            item.transformation.value,
            item.star_name.value,
            "主星" if item.star_category.value == "major" else "輔星",
            item.earthly_branch,
            item.palace_name.value,
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(table, (2.2, 2.6, 2.2, 2.0, 3.0))
    _format_table(table, header=True)


def _add_natal_interpretation(
    document: Document,
    interpretation: InterpretationResult,
) -> None:
    document.add_paragraph("本命解讀", style="Heading 1")
    _add_ai_model_metadata(document, interpretation)
    document.add_paragraph("總覽", style="Heading 2")
    document.add_paragraph(interpretation.overview)
    document.add_paragraph("十二宮解讀", style="Heading 2")
    for item in interpretation.palace_interpretations:
        _add_fact_paragraph(document, item.palace_name.value, item.summary)
    document.add_paragraph("生年四化解讀", style="Heading 2")
    document.add_paragraph(interpretation.transformation_analysis)
    document.add_paragraph("整體分析", style="Heading 2")
    _add_analysis_sections(
        document,
        (
            ("性格", interpretation.overall.personality),
            ("職涯", interpretation.overall.career),
            ("財務", interpretation.overall.finance),
            ("感情", interpretation.overall.relationships),
            ("人際", interpretation.overall.interpersonal),
            ("家庭", interpretation.overall.family),
            ("優勢", interpretation.overall.strengths),
            ("可留意的挑戰", interpretation.overall.potential_challenges),
        ),
    )


def _add_major_luck_deterministic(document: Document, chart: BasicChartResult) -> None:
    result = chart.major_luck
    gender_text = "女性" if result.gender.value == "female" else "男性"
    gender_marker = "女" if result.gender.value == "female" else "男"
    document.add_page_break()
    document.add_paragraph("大限基本資訊", style="Heading 1")
    _add_key_value_table(
        document,
        (
            ("出生年干", result.year_heavenly_stem),
            ("年干陰陽", result.year_yinyang.value),
            ("性別", gender_text),
            ("陰陽性別", f"{result.year_yinyang.value}{gender_marker}"),
            ("大限方向", result.direction.value),
            ("五行局", result.bureau_name),
            ("起限虛歲", str(result.periods[0].start_nominal_age)),
        ),
    )
    document.add_paragraph("十二大限表格", style="Heading 1")
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("大限序號", "歲數", "地支", "宮位", "宮干支", "大限四化"), strict=True):
        cell.text = text
    groups = {item.major_luck_index: item for item in result.period_transformations}
    for period in result.periods:
        group = groups[period.index]
        changes = "\n".join(
            f"{item.transformation_type.value.replace('化', '')} {item.star_name.value}"
            for item in group.transformations
        )
        cells = table.add_row().cells
        values = (
            str(period.index),
            f"{period.start_nominal_age}–{period.end_nominal_age}",
            period.earthly_branch,
            period.palace_name.value,
            period.palace_ganzhi.display,
            changes,
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(table, (1.5, 2.2, 1.2, 2.2, 1.8, 5.5))
    _format_table(table, header=True, font_size=9.5)


def _flow_year_active_major_luck_text(flow_year: FlowYearResult) -> str:
    period = flow_year.active_major_luck
    if period is not None:
        return (
            f"{period.start_nominal_age}–{period.end_nominal_age}｜"
            f"{period.palace_ganzhi.display}｜{period.palace_name.value}"
        )
    if flow_year.before_first_major_luck:
        return "尚未進入第一大限"
    return "超出目前支援的大限範圍"


def _add_flow_date_target(document: Document, flow_date: FlowDateResult) -> None:
    document.add_page_break()
    document.add_paragraph("運限查詢時間", style="Heading 1")
    solar = flow_date.target_solar_datetime.strftime("%Y-%m-%d %H:%M")
    lunar = _lunar_text(
        flow_date.target_lunar_year,
        flow_date.target_lunar_month,
        flow_date.target_lunar_day,
        flow_date.target_is_leap_month,
    )
    _add_key_value_table(
        document,
        (
            ("目標國曆時間", solar),
            ("目標農曆日期", lunar),
            ("目標有效月份", str(flow_date.target_effective_month)),
        ),
    )


def _add_flow_query_target(document: Document, result: FlowQueryResult) -> None:
    document.add_page_break()
    document.add_paragraph("運限查詢", style="Heading 1")
    target = result.normalized_target
    if isinstance(result.original_query, SolarFlowQueryInput):
        if target.original_solar_date is None or target.lunar_month is None or target.lunar_day is None:
            raise ValueError("Gregorian Flow-Query report requires complete normalized dates")
        rows = (
            ("查詢類型", "國曆"),
            ("輸入日期", target.original_solar_date.isoformat()),
            (
                "轉換農曆",
                _lunar_text(
                    target.lunar_year,
                    target.lunar_month,
                    target.lunar_day,
                    target.is_leap_month,
                ),
            ),
            ("有效月份", str(target.effective_month)),
        )
    else:
        query = result.original_query
        if not isinstance(query, LunarFlowQueryInput):
            raise ValueError("Flow-Query report requires a canonical original query")
        original = str(query.year)
        if query.month is not None:
            original += f" {'閏' if query.is_leap_month else ''}{_month_text(query.month)}月"
        if query.day is not None:
            original += _day_text(query.day)
        lunar_rows: list[tuple[str, str]] = [("查詢類型", "農曆"), ("輸入", original)]
        if target.converted_solar_date is not None:
            lunar_rows.append(("換算國曆", target.converted_solar_date.isoformat()))
        if target.effective_month is not None:
            lunar_rows.append(("有效月份", str(target.effective_month)))
        rows = tuple(lunar_rows)
    _add_key_value_table(document, rows)


def _add_flow_month_deterministic(document: Document, result: FlowMonthResult) -> None:
    document.add_page_break()
    document.add_paragraph("流月基本資訊", style="Heading 1")
    _add_key_value_table(
        document,
        (
            ("農曆年", str(result.lunar_year)),
            ("農曆月", f"{'閏' if result.is_leap_month else ''}{result.lunar_month} 月"),
            ("有效月份", str(result.effective_month)),
            ("流月干支", result.month_ganzhi.display),
            ("流月命宮", result.flow_month_life_palace_branch),
            ("所落本命宮", result.natal_host_palace_name.value),
            ("本命宮干支", result.natal_host_palace_ganzhi.display),
        ),
    )
    document.add_paragraph("流月十二宮表格", style="Heading 1")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("流月宮位", "地支", "所落本命宮", "本命宮干支"), strict=True):
        cell.text = text
    for palace in result.palaces:
        cells = table.add_row().cells
        values = (
            palace.flow_palace_name.value,
            palace.earthly_branch,
            palace.natal_palace_name.value,
            palace.natal_palace_ganzhi.display,
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(table, (3.2, 2.2, 3.8, 3.2))
    _format_table(table, header=True)

    document.add_paragraph("流月四化", style="Heading 1")
    transformations = document.add_table(rows=1, cols=5)
    transformations.style = "Table Grid"
    for cell, text in zip(transformations.rows[0].cells, ("四化", "星曜", "類型", "所落本命宮", "地支"), strict=True):
        cell.text = text
    for item in result.transformations:
        cells = transformations.add_row().cells
        values = (
            item.transformation_type.value,
            item.star_name.value,
            "主星" if item.star_category.value == "major" else "輔星",
            item.natal_palace_name.value,
            item.natal_branch,
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(transformations, (2.0, 2.6, 2.2, 3.8, 2.0))
    _format_table(transformations, header=True)


def _add_flow_day_deterministic(document: Document, result: FlowDayResult) -> None:
    document.add_page_break()
    document.add_paragraph("流日基本資訊", style="Heading 1")
    _add_key_value_table(
        document,
        (
            (
                "目標農曆日期",
                _lunar_text(
                    result.target_lunar_year,
                    result.target_lunar_month,
                    result.target_lunar_day,
                    result.target_is_leap_month,
                ),
            ),
            ("流日干支", result.day_ganzhi.display),
            ("流日命宮", result.flow_day_life_palace_branch),
            ("所落本命宮", result.natal_host_palace_name.value),
            ("本命宮干支", result.natal_host_palace_ganzhi.display),
        ),
    )
    document.add_paragraph("流日十二宮表格", style="Heading 1")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("流日宮位", "地支", "所落本命宮", "本命宮干支"), strict=True):
        cell.text = text
    for palace in result.palaces:
        cells = table.add_row().cells
        values = (
            palace.flow_palace_name.value,
            palace.earthly_branch,
            palace.natal_palace_name.value,
            palace.natal_palace_ganzhi.display,
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(table, (3.2, 2.2, 3.8, 3.2))
    _format_table(table, header=True)

    document.add_paragraph("流日四化", style="Heading 1")
    transformations = document.add_table(rows=1, cols=5)
    transformations.style = "Table Grid"
    for cell, text in zip(transformations.rows[0].cells, ("四化", "星曜", "類型", "所落本命宮", "地支"), strict=True):
        cell.text = text
    for item in result.transformations:
        cells = transformations.add_row().cells
        values = (
            item.transformation_type.value,
            item.star_name.value,
            "主星" if item.star_category.value == "major" else "輔星",
            item.natal_palace_name.value,
            item.natal_branch,
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(transformations, (2.0, 2.6, 2.2, 3.8, 2.0))
    _format_table(transformations, header=True)


def _add_flow_period_transformation_analysis(document: Document, heading: str, interpretation) -> None:
    document.add_paragraph(heading, style="Heading 2")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("四化", "星曜", "本命宮位", "解讀"), strict=True):
        cell.text = text
    for item in interpretation.transformation_analysis:
        cells = table.add_row().cells
        for cell, text in zip(
            cells,
            (item.transformation_type.value, item.star_name.value, item.natal_palace_name.value, item.analysis),
            strict=True,
        ):
            cell.text = text
    _format_table(table, header=True)


def _add_flow_month_interpretation(
    document: Document, interpretation: FlowMonthInterpretationResult
) -> None:
    document.add_page_break()
    document.add_paragraph("流月解讀", style="Heading 1")
    _add_ai_model_metadata(document, interpretation)
    _add_key_value_table(
        document,
        (
            ("農曆年月", f"{interpretation.lunar_year} {'閏' if interpretation.is_leap_month else ''}{interpretation.lunar_month} 月"),
            ("流月干支", interpretation.month_ganzhi.display),
            ("流月命宮", interpretation.flow_month_life_palace_branch),
            ("所落本命宮", interpretation.natal_host_palace.value),
        ),
    )
    _add_analysis_sections(document, (
        ("流月總覽", interpretation.overview),
        ("流月命宮分析", interpretation.life_palace_analysis),
        ("大限背景", interpretation.major_luck_context),
        ("流年背景", interpretation.flow_year_context),
    ))
    _add_flow_period_transformation_analysis(document, "流月四化分析", interpretation)
    _add_analysis_sections(document, (
        ("職涯", interpretation.career),
        ("財務", interpretation.finance),
        ("感情", interpretation.relationships),
        ("家庭／人際", interpretation.family_and_interpersonal),
        ("優勢", interpretation.strengths),
        ("可留意的挑戰", interpretation.potential_challenges),
        ("實際重點", interpretation.practical_focus),
    ))


def _add_flow_day_interpretation(
    document: Document, interpretation: FlowDayInterpretationResult
) -> None:
    document.add_page_break()
    document.add_paragraph("流日解讀", style="Heading 1")
    _add_ai_model_metadata(document, interpretation)
    _add_key_value_table(
        document,
        (
            ("農曆日期", _lunar_text(
                interpretation.lunar_year,
                interpretation.lunar_month,
                interpretation.lunar_day,
                interpretation.is_leap_month,
            )),
            ("流日干支", interpretation.day_ganzhi.display),
            ("流日命宮", interpretation.flow_day_life_palace_branch),
            ("所落本命宮", interpretation.natal_host_palace.value),
        ),
    )
    _add_analysis_sections(document, (
        ("流日總覽", interpretation.overview),
        ("流日命宮分析", interpretation.life_palace_analysis),
        ("大限背景", interpretation.major_luck_context),
        ("流年背景", interpretation.flow_year_context),
        ("流月背景", interpretation.flow_month_context),
    ))
    _add_flow_period_transformation_analysis(document, "流日四化分析", interpretation)
    _add_analysis_sections(document, (
        ("工作", interpretation.work),
        ("財務", interpretation.finance),
        ("感情", interpretation.relationships),
        ("家庭／人際", interpretation.family_and_interpersonal),
        ("優勢", interpretation.strengths),
        ("可留意的挑戰", interpretation.potential_challenges),
        ("實際重點", interpretation.practical_focus),
    ))


def _add_flow_year_deterministic(document: Document, flow_year: FlowYearResult) -> None:
    life_host = next(item for item in flow_year.palaces if item.flow_palace_name is PalaceName.LIFE)
    document.add_page_break()
    document.add_paragraph("流年基本資訊", style="Heading 1")
    _add_key_value_table(
        document,
        (
            ("流年年份", str(flow_year.target_lunar_year)),
            ("流年干支", flow_year.ganzhi.display),
            ("虛歲", str(flow_year.nominal_age)),
            ("流年命宮", flow_year.flow_life_palace_branch),
            ("所落本命宮", life_host.natal_palace_name.value),
            ("本命宮干支", life_host.natal_palace_ganzhi.display),
            ("目前大限", _flow_year_active_major_luck_text(flow_year)),
        ),
    )
    document.add_paragraph("流年十二宮表格", style="Heading 1")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("流年宮位", "地支", "所落本命宮", "本命宮干支"), strict=True):
        cell.text = text
    for palace in flow_year.palaces:
        cells = table.add_row().cells
        values = (
            palace.flow_palace_name.value,
            palace.earthly_branch,
            palace.natal_palace_name.value,
            palace.natal_palace_ganzhi.display,
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(table, (3.2, 2.2, 3.8, 3.2))
    _format_table(table, header=True)

    document.add_paragraph("流年四化", style="Heading 1")
    transformations = document.add_table(rows=1, cols=5)
    transformations.style = "Table Grid"
    for cell, text in zip(transformations.rows[0].cells, ("四化", "星曜", "類型", "所落本命宮", "地支"), strict=True):
        cell.text = text
    for item in flow_year.transformations:
        cells = transformations.add_row().cells
        values = (
            item.transformation_type.value,
            item.star_name.value,
            "主星" if item.star_category.value == "major" else "輔星",
            item.natal_palace_name.value,
            item.earthly_branch,
        )
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _set_table_widths(transformations, (2.0, 2.6, 2.2, 3.8, 2.0))
    _format_table(transformations, header=True)


def _add_major_luck_interpretation(
    document: Document,
    interpretation: MajorLuckInterpretationResult,
) -> None:
    document.add_page_break()
    document.add_paragraph("大限解讀", style="Heading 1")
    _add_ai_model_metadata(document, interpretation)
    _add_key_value_table(
        document,
        (
            ("大限序號", str(interpretation.major_luck_index)),
            ("歲數範圍", f"{interpretation.start_nominal_age}–{interpretation.end_nominal_age}"),
            ("宮位", interpretation.palace_name.value),
            ("地支", interpretation.earthly_branch),
            ("干支", interpretation.palace_ganzhi.display),
        ),
    )
    _add_analysis_sections(
        document,
        (
            ("大限總覽", interpretation.overview),
            ("大限所在宮分析", interpretation.host_palace_analysis),
        ),
    )
    document.add_paragraph("大限四化分析", style="Heading 2")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("四化", "星曜", "本命宮位", "解讀"), strict=True):
        cell.text = text
    for item in interpretation.transformation_analysis:
        cells = table.add_row().cells
        for cell, text in zip(
            cells,
            (item.transformation_type.value, item.star_name.value, item.natal_palace_name.value, item.analysis),
            strict=True,
        ):
            cell.text = text
    _format_table(table, header=True)
    _add_analysis_sections(
        document,
        (
            ("職涯", interpretation.career),
            ("財務", interpretation.finance),
            ("感情", interpretation.relationships),
            ("家庭／人際", interpretation.family_and_interpersonal),
            ("優勢", interpretation.strengths),
            ("可留意的挑戰", interpretation.potential_challenges),
            ("實際重點", interpretation.practical_focus),
        ),
    )


def _flow_year_major_luck_text(interpretation: FlowYearInterpretationResult) -> str:
    anchor = interpretation.active_major_luck_summary
    if anchor.status is ActiveMajorLuckStatus.BEFORE_FIRST:
        return "尚未進入第一大限"
    if anchor.status is ActiveMajorLuckStatus.AFTER_SUPPORTED:
        return "超出目前支援的大限範圍"
    return (
        f"{anchor.start_nominal_age}–{anchor.end_nominal_age}｜"
        f"{anchor.palace_ganzhi.display}｜{anchor.palace_name.value}"
    )


def _add_flow_year_interpretation(
    document: Document,
    interpretation: FlowYearInterpretationResult,
) -> None:
    document.add_page_break()
    document.add_paragraph("流年解讀", style="Heading 1")
    _add_ai_model_metadata(document, interpretation)
    host = interpretation.flow_life_palace_natal_host
    _add_key_value_table(
        document,
        (
            ("流年年份", str(interpretation.target_year)),
            ("流年干支", interpretation.flow_year_ganzhi.display),
            ("虛歲", str(interpretation.nominal_age)),
            ("流年命宮", interpretation.flow_life_palace_branch),
            ("所落本命宮", f"{host.palace_name.value}｜{host.palace_ganzhi.display}"),
            ("目前大限", _flow_year_major_luck_text(interpretation)),
        ),
    )
    _add_analysis_sections(
        document,
        (
            ("流年總覽", interpretation.overview),
            ("流年命宮分析", interpretation.flow_life_palace_analysis),
            ("大限背景", interpretation.major_luck_context),
        ),
    )
    document.add_paragraph("流年四化分析", style="Heading 2")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("四化", "星曜", "本命宮位", "解讀"), strict=True):
        cell.text = text
    for item in interpretation.transformation_analysis:
        cells = table.add_row().cells
        for cell, text in zip(
            cells,
            (item.transformation_type.value, item.star_name.value, item.natal_palace_name.value, item.analysis),
            strict=True,
        ):
            cell.text = text
    _format_table(table, header=True)
    _add_analysis_sections(
        document,
        (
            ("職涯", interpretation.career),
            ("財務", interpretation.finance),
            ("感情", interpretation.relationships),
            ("家庭／人際", interpretation.family_and_interpersonal),
            ("優勢", interpretation.strengths),
            ("可留意的挑戰", interpretation.potential_challenges),
            ("實際重點", interpretation.practical_focus),
        ),
    )


def render_word_report(
    chart: BasicChartResult,
    birth_input: BirthInput,
    interpretation: InterpretationResult,
    *,
    flow_date: FlowDateResult | None = None,
    flow_query: FlowQueryResult | None = None,
    flow_year: FlowYearResult | None = None,
    major_luck_interpretation: MajorLuckInterpretationResult | None = None,
    flow_year_interpretation: FlowYearInterpretationResult | None = None,
    flow_month_interpretation: FlowMonthInterpretationResult | None = None,
    flow_day_interpretation: FlowDayInterpretationResult | None = None,
    output_dir: Path | str = DEFAULT_REPORT_DIR,
    chart_output_dir: Path | str | None = None,
) -> RenderedWordReport:
    """Generate one DOCX without invoking any LLM service."""

    if flow_date is not None and flow_query is not None:
        raise ValueError("use either legacy Flow-Date or Flow-Query facts")
    if flow_query is not None:
        if flow_year is None:
            flow_year = flow_query.flow_year
        elif flow_year != flow_query.flow_year:
            raise ValueError("Flow-Query and Flow-Year facts do not match")
    if flow_date is not None:
        if flow_year is None:
            flow_year = flow_date.flow_year
        elif flow_year != flow_date.flow_year:
            raise ValueError("Flow-Date and Flow-Year facts do not match")
    if flow_year_interpretation is not None and flow_year is None:
        raise ValueError("Flow-Year interpretation requires deterministic Flow-Year facts")
    if (
        flow_year_interpretation is not None
        and flow_year_interpretation.target_year != flow_year.target_lunar_year
    ):
        raise ValueError("Flow-Year interpretation does not match deterministic Flow-Year facts")
    if flow_month_interpretation is not None and (flow_query is None or flow_query.flow_month is None):
        raise ValueError("Flow-Month interpretation requires deterministic Flow-Month facts")
    if flow_day_interpretation is not None and (flow_query is None or flow_query.flow_day is None):
        raise ValueError("Flow-Day interpretation requires deterministic Flow-Day facts")

    report_dir = Path(output_dir)
    report_dir.mkdir(parents=True, exist_ok=True)
    chart_render = render_chart_png(chart, chart_output_dir or report_dir / "charts")

    document = Document()
    _configure_document(document)
    document.core_properties.title = "Tiger-ZiWei 紫微斗數完整報告"
    document.core_properties.subject = "紫微斗數完整命盤與結構化解讀"

    title = document.add_paragraph("Tiger-ZiWei 紫微斗數完整報告", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("出生資料", style="Heading 1")
    lunar = chart.calendar.lunar_date
    lunar_text = _lunar_text(lunar.year, lunar.month, lunar.day, lunar.is_leap_month)
    input_mode = "農曆" if birth_input.calendar_type is CalendarType.LUNAR else "國曆"
    birth_rows = (
        ("姓名", chart.birth_data.name or "未填寫"),
        ("性別", "女" if chart.birth_data.gender.value == "female" else "男"),
        ("輸入方式", input_mode),
        ("原始出生日期輸入", _original_input_text(birth_input)),
        ("換算西元日期", chart.calendar.solar_date.isoformat()),
        ("農曆日期", lunar_text),
        ("是否閏月", "是" if lunar.is_leap_month else "否"),
        ("出生時間", f"{chart.birth_data.birth_hour:02d}:{chart.birth_data.birth_minute:02d}"),
        ("出生地", chart.birth_data.birthplace),
    )
    _add_key_value_table(document, birth_rows)

    document.add_page_break()
    chart_heading = document.add_paragraph("命盤圖片", style="Heading 1")
    chart_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(0)
    image_paragraph.add_run().add_picture(str(chart_render.path), width=Inches(6.65))

    document.add_page_break()
    document.add_paragraph("本命基本資訊", style="Heading 1")
    _add_key_value_table(
        document,
        (
            ("農曆日期", lunar_text),
            ("年柱", chart.year_ganzhi.display),
            ("月柱", chart.month_ganzhi.display),
            ("日柱", chart.day_ganzhi.display),
            ("時柱", chart.hour_ganzhi.display),
            ("命宮", chart.life_palace_branch),
            ("身宮", chart.body_palace_branch),
            ("身宮宿宮", chart.body_palace_name.value),
            ("五行局", chart.five_elements_bureau.bureau_name),
            ("出生年干", chart.year_ganzhi.heavenly_stem),
            ("實際農曆月", str(lunar.month)),
            ("是否閏月", "是" if lunar.is_leap_month else "否"),
            ("有效本命月", str(chart.palace_layout.effective_lunar_month)),
        ),
    )
    document.add_paragraph("採用規則", style="Heading 2")
    for rule in ADOPTED_RULES:
        document.add_paragraph(rule, style="List Bullet")

    _add_natal_palace_table(document, chart)
    _add_birth_year_transformations(document, chart)
    _add_natal_interpretation(document, interpretation)
    _add_major_luck_deterministic(document, chart)

    if major_luck_interpretation is not None:
        _add_major_luck_interpretation(document, major_luck_interpretation)
    if flow_date is not None:
        _add_flow_date_target(document, flow_date)
    if flow_query is not None:
        _add_flow_query_target(document, flow_query)
    if flow_year is not None:
        _add_flow_year_deterministic(document, flow_year)
    if flow_year_interpretation is not None:
        _add_flow_year_interpretation(document, flow_year_interpretation)
    if flow_date is not None:
        _add_flow_month_deterministic(document, flow_date.flow_month)
        _add_flow_day_deterministic(document, flow_date.flow_day)
    if flow_query is not None and flow_query.flow_month is not None:
        _add_flow_month_deterministic(document, flow_query.flow_month)
        if flow_month_interpretation is not None:
            _add_flow_month_interpretation(document, flow_month_interpretation)
    if flow_query is not None and flow_query.flow_day is not None:
        _add_flow_day_deterministic(document, flow_query.flow_day)
        if flow_day_interpretation is not None:
            _add_flow_day_interpretation(document, flow_day_interpretation)

    document.add_paragraph("免責說明", style="Heading 1")
    document.add_paragraph(DISCLAIMER)

    filename = _safe_filename(chart, flow_date, flow_query)
    destination = report_dir / filename
    document.save(destination)
    return RenderedWordReport(path=destination, filename=filename)
